"""The Jobs board as a formatted .xlsx (opens in Google Sheets with its table, colours and Status dropdown)."""

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

import owner

from .pipeline import platform

TZ = ZoneInfo(owner.OWNER.get("timezone") or "UTC")
# status -> (label, fill colour); the label list is also the Status dropdown
STATUS = {
    "drafted": ("To approve", "DCE7F7"), "needs_you": ("Needs you", "FCE3C7"), "draft_blocked": ("Needs you", "FCE3C7"),
    "handoff": ("Finish on PC", "FFF1B8"), "queued": ("Queued", "ECECEC"), "new": ("Queued", "ECECEC"),
    "approved": ("Applying", "E3DCF7"), "applying": ("Applying", "E3DCF7"), "applied": ("Applied", "D6F0DC"),
    "follow_up": ("Follow up", "FFD9D9"), "replied": ("Replied", "B9E6C5"), "closed": ("Closed", "E0E0E0"),
    "skipped": ("Skipped", "F2F2F2"),
}
COLS = [("Platform", 13), ("Job title", 46), ("Company", 24), ("Status", 15), ("Score", 8), ("Pay", 22),
        ("Found", 17), ("Applied", 17), ("Followed up", 17), ("Replied", 17), ("Notes", 48), ("Link", 9)]


def _when(iso: str):
    try:
        return datetime.fromisoformat(iso).astimezone(TZ).replace(tzinfo=None) if iso else None
    except ValueError:
        return None


def _status(j: dict) -> str:
    key = "follow_up" if j.get("status") == "applied" and j.get("follow_up") == "due" else j.get("status", "")
    return STATUS.get(key, (key.replace("_", " ").capitalize(), ""))[0]


def build(jobs: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append([c for c, _ in COLS])
    for j in jobs:
        company = "" if j.get("company") in ("OnlineJobs.ph", "Indeed") else j.get("company", "")
        ws.append([platform(j.get("url")), j.get("title", ""), company, _status(j), j.get("score"), j.get("pay", ""),
                   _when(j.get("found")), _when(j.get("applied_at")), _when(j.get("followed_up_at")),
                   _when(j.get("replied_at")), {"submitted": "Application sent", "closed": "Job closed"}.get(j.get("apply_result"), j.get("apply_result") or j.get("needs_why") or ""), "Open"])
        link = ws.cell(ws.max_row, len(COLS))
        link.hyperlink, link.font = j.get("url", ""), Font(color="1F5FBF", underline="single")
    last = max(ws.max_row, 2)

    for i, (_, width) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.alignment = Alignment(vertical="center", wrap_text=c.column in (2, 11))
        for c in row[6:10]:
            c.number_format = "d mmm, h:mm AM/PM"
        row[4].number_format = "0.0"
        row[4].alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 24

    ref = f"A1:{get_column_letter(len(COLS))}{last}"
    table = Table(displayName="Jobs", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    labels = list(dict.fromkeys(v[0] for v in STATUS.values()))
    dv = DataValidation(type="list", formula1='"' + ",".join(labels) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D2:D{last}")
    for label, colour in dict(STATUS.values()).items():  # colours follow the dropdown when he changes a status
        ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="equal", formula=[f'"{label}"'],
                                      fill=PatternFill("solid", start_color=colour), font=Font(bold=True)))
    ws.conditional_formatting.add(f"E2:E{last}", ColorScaleRule(start_type="num", start_value=1, start_color="F8D7D7",
                                  mid_type="num", mid_value=3, mid_color="FFF4CC", end_type="num", end_value=5, end_color="CDEFD6"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    from openpyxl import load_workbook
    data = build([{"url": "https://www.reddit.com/r/forhire/x", "title": "Editor", "company": "Reddit /u/x", "status": "applied",
                   "follow_up": "due", "score": 4.0, "found": "2026-10-27T00:00:00+00:00"}])
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A2"].value == "Reddit" and ws["D2"].value == "Follow up" and ws["L2"].hyperlink.target.startswith("https://")
    assert ws["G2"].value == datetime(2026, 10, 27, 8, 0) or owner.OWNER.get("timezone") != "Asia/Manila"
    print("export_xlsx selftest ok")
